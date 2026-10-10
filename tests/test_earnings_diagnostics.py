"""Known accounting examples, missing evidence and incompatible vintages."""
from copy import deepcopy
from decimal import Decimal
import io

import pytest
from openpyxl import load_workbook

from src.harness.diagnostics import comparison_groups, earnings_bridge, save_diagnostics
from src.harness.store import Run


def rows():
    periods = {"2024-12-31": {"revenue": 100, "cost_of_revenue": 60, "pretax_profit": 30,
                            "income_tax_expense": 5, "net_profit": 25, "net_profit_parent": 20,
                            "operating_cash_flow": 10, "capex_cash_paid": 15},
               "2025-12-31": {"revenue": 150, "cost_of_revenue": 105, "pretax_profit": 40,
                            "income_tax_expense": 8, "net_profit": 32, "net_profit_parent": 26,
                            "operating_cash_flow": 40, "capex_cash_paid": 20}}
    return [{"stock_code": "000000", "company": "合成公司", "statement_scope": "consolidated",
             "published_at": "2026-03-01", "revision_flag": "original", "source_database": "synthetic",
             "period_type": "annual", "report_period": period, "canonical_metric": metric,
             "value_status": "observed", "value_cny": str(amount), "source_file": "synthetic.txt",
             "source_locator": "line:1", "source_sha256": "0" * 64, "unit": "元", "value_original": str(amount)}
            for period, metrics in periods.items() for metric, amount in metrics.items()]


def diagnose(data=None):
    return earnings_bridge(rows() if data is None else data, "2024-12-31", "2025-12-31")


def test_known_profit_bridge_cash_denominator_and_provenance():
    result = diagnose()
    assert result["bridge_target"] == "归母净利润"
    assert [Decimal(r["贡献_元"]) for r in result["bridge"]] == [20, -15, 5, -3, -1]
    assert Decimal(result["target_change_cny"]) == Decimal(result["bridge_sum_cny"]) == 6
    indicators = {r["id"]: r for r in result["indicators"]}
    assert Decimal(indicators["cash_conversion"]["前期"]) == Decimal("0.4")  # Uses consolidated, not parent profit.
    assert Decimal(indicators["ocf_less_capex"]["前期"]) == -5
    assert all(row["sources"] for row in result["bridge"])
    assert any("毛利率下降" in f for f in result["findings"])
    assert any("FCFF" in f for f in result["limitations"])


def test_missing_profit_fields_partial_bridge_does_not_fill_zero():
    result = diagnose([r for r in rows() if r["canonical_metric"] not in {"net_profit", "income_tax_expense"}])
    assert result["bridge_target"] == "毛利" and len(result["bridge"]) == 2
    assert result["missing_fields"] and result["warnings"]
    assert "cash_conversion" not in {r["id"] for r in result["indicators"]}
    assert Decimal(result["bridge_sum_cny"]) == 5


@pytest.mark.parametrize("field,value", [("stock_code", "000001"), ("statement_scope", "parent"),
                                        ("revision_flag", "restated"), ("published_at", "2025-03-01"),
                                        ("source_database", "other")])
def test_incompatible_company_scope_revision_and_disclosure_vintage_rejected(field, value):
    data = rows()
    data[0][field] = value
    with pytest.raises(ValueError, match="同公司"):
        diagnose(data)


def test_duplicate_metric_rejected_even_equal_value():
    data = rows()
    data.append(deepcopy(data[0]))
    with pytest.raises(ValueError, match="多条"):
        diagnose(data)


def test_missing_and_zero_revenue_rejected():
    with pytest.raises(ValueError, match="缺失值"):
        diagnose([r for r in rows() if r["canonical_metric"] != "revenue"])
    data = rows()
    data[0]["value_cny"] = "0"
    with pytest.raises(ValueError, match="正数"):
        diagnose(data)


def test_profit_identity_mismatch_rejected():
    data = rows()
    next(r for r in data if r["canonical_metric"] == "net_profit")["value_cny"] = "999"
    with pytest.raises(ValueError, match="勾稽"):
        diagnose(data)


def test_interim_and_nonconsecutive_years_not_treated_as_annual():
    with pytest.raises(ValueError, match="自然年度"):
        earnings_bridge(rows(), "2024-06-30", "2025-06-30")
    with pytest.raises(ValueError, match="连续两年"):
        earnings_bridge(rows(), "2023-12-31", "2025-12-31")
    data = rows()
    for r in data:
        r["period_type"] = "cumulative"
    assert comparison_groups(data) == []


def test_negative_profit_cash_ratio_omitted_not_reinterpreted():
    data = rows()
    for r in data:
        if r["report_period"] == "2024-12-31":
            r["value_cny"] = {"net_profit": "-5", "pretax_profit": "0", "net_profit_parent": "-10"}.get(r["canonical_metric"], r["value_cny"])
    result = diagnose(data)
    assert "cash_conversion" not in {r["id"] for r in result["indicators"]}
    assert any("非正" in warning for warning in result["warnings"])


def test_export_keeps_excel_source_literal_and_reconciled_results(tmp_path):
    data = rows()
    for r in data:
        r["source_locator"] = '=HYPERLINK("unsafe")'
    run = Run(tmp_path, "diagnostics", [])
    result = diagnose(data)
    save_diagnostics(run, result)
    workbook = load_workbook(io.BytesIO((run.path / "outputs/earnings_diagnostics.xlsx").read_bytes()))
    assert {"两期指标", "盈利变化桥", "原件定位"} == set(workbook.sheetnames)
    assert not any(c.data_type == "f" for sheet in workbook for row in sheet for c in row)
    assert "会计分解" in (run.path / "outputs/earnings_diagnostics.md").read_text(encoding="utf-8")
    assert run.metadata["status"] == "earnings_diagnostics_completed"
