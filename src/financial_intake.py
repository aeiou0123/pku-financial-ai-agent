"""Offline financial evidence intake; never guesses vendor field codes or updates forecasts."""
from __future__ import annotations

import csv
import hashlib
import json
import re
from zipfile import BadZipFile
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

FIELDS = [
    "stock_code", "company", "report_period", "published_at", "statement_scope",
    "period_type", "metric_original", "canonical_metric", "value_original", "unit",
    "currency", "revision_flag", "source_database", "source_file", "source_locator",
    "source_sha256",
]
STOCK_METRICS = {
    "total_assets", "total_liabilities", "total_equity", "parent_equity", "cash",
    "accounts_receivable", "inventory", "accounts_payable", "fixed_assets",
    "short_term_debt", "long_term_debt", "current_debt_maturities",
}
FLOW_METRICS = {
    "revenue", "cost_of_revenue", "net_profit", "net_profit_parent", "pretax_profit",
    "income_tax_expense", "selling_expense", "admin_expense", "rd_expense",
    "operating_cash_flow", "capex_cash_paid",
}
UNITS = {"元": Decimal(1), "万元": Decimal(10000), "亿元": Decimal(100000000),
         "cny": Decimal(1), "bn_cny": Decimal(1000000000)}
MISSING = {"", "-", "—", "--", "NA", "N/A", "NULL"}


def sha256(path: Path) -> str:
    with path.open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def init_workspace(root: Path) -> list[str]:
    dirs = ["raw/A_financial", "raw/dictionaries", "raw/B_filings", "normalized", "runs"]
    for item in dirs:
        (root / item).mkdir(parents=True, exist_ok=True)
    return dirs


def inventory(root: Path) -> dict:
    """Fingerprint originals. Header previews are diagnostics, not a field mapping."""
    files = []
    for path in sorted((root / "raw").rglob("*")):
        if not path.is_file():
            continue
        if not path.resolve().is_relative_to(root.resolve()):
            files.append({"path": str(path.relative_to(root)), "error": "outside_workspace"})
            continue
        item = {"path": str(path.relative_to(root)), "bytes": path.stat().st_size,
                "sha256": sha256(path)}
        try:
            if path.suffix.lower() == ".csv":
                for encoding in ("utf-8-sig", "gb18030"):
                    try:
                        with path.open(encoding=encoding, newline="") as f:
                            item["first_rows"] = [row for _, row in zip(range(3), csv.reader(f))]
                        item["encoding"] = encoding
                        break
                    except UnicodeDecodeError:
                        continue
                else:
                    item["error"] = "unsupported_csv_encoding"
            elif path.suffix.lower() == ".xlsx":
                from openpyxl import load_workbook  # Read-only original inspection.
                workbook = load_workbook(path, read_only=True, data_only=False)
                try:
                    item["sheets"] = [{"name": s.title, "rows": s.max_row,
                                       "columns": s.max_column,
                                       "first_rows": list(s.iter_rows(min_row=1, max_row=3,
                                                                       values_only=True))}
                                      for s in workbook]
                finally:
                    workbook.close()
            elif path.suffix.lower() == ".xls":
                item["note"] = "Legacy XLS preserved; export XLSX/CSV for header inspection."
        except (ImportError, ValueError, OSError, KeyError, BadZipFile) as exc:
            item["error"] = type(exc).__name__ + ": " + str(exc)
        files.append(item)
    return {"schema": "financial_raw_inventory_v1", "files": files,
            "field_mapping": "pending_actual_vendor_dictionary"}


def _amount(text: str) -> Decimal | None:
    if text.strip().upper() in MISSING:
        return None
    # Permit plain decimals or correctly grouped thousands; reject locale ambiguity.
    if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?", text.strip()):
        raise ValueError("invalid_number")
    try:
        value = Decimal(text.strip().replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError("invalid_number") from exc
    if not value.is_finite():
        raise ValueError("nonfinite_number")
    return value


def validate(input_path: Path, source_root: Path, as_of: date) -> tuple[list[dict], dict]:
    with input_path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != FIELDS:
            raise ValueError("Header must exactly match financial_long_v2.csv")
        rows = list(reader)
    errors, warnings, normalized = [], [], []
    hashes = {}
    for number, row in enumerate(rows, 2):
        try:
            if None in row or any(v is None for v in row.values()):
                raise ValueError("malformed_csv_row")
            if not re.fullmatch(r"[0-9]{6}", row["stock_code"]):
                raise ValueError("stock_code_requires_six_digits")
            for key in ("company", "metric_original", "revision_flag", "source_database",
                        "source_locator", "source_file"):
                if not row[key].strip():
                    raise ValueError("missing_" + key)
            for key in ("report_period", "published_at"):
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", row[key]):
                    raise ValueError("invalid_" + key)
            period, published = date.fromisoformat(row["report_period"]), date.fromisoformat(row["published_at"])
            if not period <= published <= as_of:
                raise ValueError("publication_or_period_after_cutoff")
            if row["statement_scope"] not in {"consolidated", "parent"}:
                raise ValueError("unknown_statement_scope")
            metric = row["canonical_metric"]
            if metric not in STOCK_METRICS | FLOW_METRICS:
                raise ValueError("unknown_canonical_metric")
            kind = row["period_type"]
            if metric in STOCK_METRICS and kind != "point_in_time":
                raise ValueError("stock_metric_requires_point_in_time")
            if metric in FLOW_METRICS and kind not in {"annual", "cumulative", "single_quarter"}:
                raise ValueError("invalid_flow_period_type")
            if kind == "annual" and (period.month, period.day) != (12, 31):
                raise ValueError("annual_flow_requires_december_31")
            if row["unit"] not in UNITS or row["currency"] != "CNY":
                raise ValueError("unsupported_unit_or_currency")
            relative = Path(row["source_file"])
            if relative.is_absolute():
                raise ValueError("source_file_requires_relative_path")
            source = (source_root / relative).resolve()
            if not source.is_relative_to(source_root.resolve()):
                raise ValueError("source_file_outside_root")
            if not source.is_file():
                raise ValueError("source_file_missing")
            if source not in hashes:
                hashes[source] = sha256(source)
            if row["source_sha256"] != hashes[source]:
                raise ValueError("source_hash_mismatch")
            amount = _amount(row["value_original"])
            cny = None if amount is None else amount * UNITS[row["unit"]]
            normalized.append({**row, "input_line": number,
                               "value_cny": "" if cny is None else str(cny),
                               "value_bn_cny": "" if cny is None else str(cny / Decimal(10**9)),
                               "value_status": "missing" if cny is None else "observed"})
            if cny is None:
                warnings.append({"line": number, "issue": "missing_value_not_zero"})
        except (ValueError, OSError) as exc:
            errors.append({"line": number, "issue": str(exc)})

    # Never select one duplicate or revision silently. Group each source/version separately.
    groups = defaultdict(list)
    for row in normalized:
        key = tuple(row[k] for k in ("stock_code", "report_period", "statement_scope",
                                    "published_at", "revision_flag", "source_database"))
        groups[key].append(row)
    checks = []
    for key, records in groups.items():
        index = defaultdict(list)
        for row in records:
            index[(row["canonical_metric"], row["period_type"])].append(row)
        duplicates = [k for k, v in index.items() if len(v) > 1]
        group = dict(zip(("stock_code", "report_period", "statement_scope", "published_at",
                          "revision_flag", "source_database"), key))
        if duplicates:
            errors.append({"group": group, "issue": "duplicate_metric_requires_review",
                           "metrics": duplicates})
            checks.append({**group, "check": "group", "status": "blocked_duplicates"})
            continue

        def get(metric, kind):
            matches = index.get((metric, kind), [])
            return None if not matches or matches[0]["value_status"] == "missing" else Decimal(matches[0]["value_cny"])

        a, liabilities, equity = [get(m, "point_in_time") for m in
                                  ("total_assets", "total_liabilities", "total_equity")]
        if any(v is None for v in (a, liabilities, equity)):
            checks.append({**group, "check": "balance_sheet", "status": "not_checked_missing_totals"})
        else:
            residual = a - liabilities - equity
            status = "passed" if abs(residual) <= Decimal("1") else "failed"
            checks.append({**group, "check": "balance_sheet", "status": status,
                           "residual_cny": str(residual), "tolerance_cny": "1"})
            if status == "failed":
                errors.append({"group": group, "issue": "balance_sheet_mismatch"})
        for kind in ("annual", "cumulative", "single_quarter"):
            if not any((m, kind) in index for m in FLOW_METRICS):
                continue
            revenue, cost = get("revenue", kind), get("cost_of_revenue", kind)
            if revenue is None or cost is None or revenue == 0:
                checks.append({**group, "check": "gross_margin", "period_type": kind,
                               "status": "not_checked_missing_or_zero_revenue"})
            else:
                checks.append({**group, "check": "gross_margin", "period_type": kind,
                               "status": "calculated", "value": str((revenue - cost) / revenue),
                               "interpretation": "historical_ratio_only"})
    complete = bool(checks) and all(c["status"] in {"passed", "calculated"} for c in checks)
    report = {"schema": "financial_quality_v1", "as_of": as_of.isoformat(),
              "input_sha256": sha256(input_path), "input_rows": len(rows),
              "normalized_rows": len(normalized), "errors": errors, "warnings": warnings,
              "checks": checks, "status": "invalid" if errors else ("checked" if complete and not warnings else "partial"),
              "source_authentication": "hash_matches_local_bytes_only; provenance_and_locator_require_human_review",
              "forecast_updated": False}
    return normalized, report


def save_run(root: Path, input_path: Path, normalized: list[dict], report: dict) -> Path:
    """Create a unique immutable run; retain even invalid input for inspection."""
    import shutil
    path = root / "runs" / (datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "_" + uuid4().hex[:8])
    path.mkdir(parents=True, exist_ok=False)
    shutil.copyfile(input_path, path / "input_snapshot.csv")
    with (path / "normalized.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS + ["input_line", "value_cny", "value_bn_cny", "value_status"])
        writer.writeheader()
        writer.writerows(normalized)
    (path / "quality_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    (path / "review.md").write_text(
        f"# 财务接收检查\n\n状态：{report['status']}。输入 {report['input_rows']} 行，标准化 {report['normalized_rows']} 行；"
        f"错误 {len(report['errors'])} 项，缺失提示 {len(report['warnings'])} 项。\n\n"
        "原值和原件指纹已保留。具体检查见 quality_report.json。此工具核对本地文件指纹和数值口径，"
        "来源可信性、原文页码和字段映射仍需人工复核；不会更新预测、估值或技术题。\n", encoding="utf-8")
    return path
