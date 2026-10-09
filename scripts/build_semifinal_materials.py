"""Build editable information and bookmarked PDF drafts from recorded evidence.

Run with the document runtime; provide the unchanged official DOCX template and
a local CJK TrueType font. Never infer team names, identities or signatures.
"""
from __future__ import annotations
import argparse
import json
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]


def make_information(template, output, content):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    document = Document(template)
    if len(document.tables) != 7 or len(document.tables[0].rows) != 4:
        raise ValueError("Unexpected official information-template structure")
    limits = [300, 200, 800, 700, 500]
    counts = [len(''.join(text.split())) for text in content['information']]
    if len(counts) != 5 or any(n > limit for n, limit in zip(counts, limits)) or sum(counts) > 2500:
        raise ValueError("Information fields exceed conservative character bounds")
    values = ['待负责人填写', 'Claim2Value', '待负责人逐人填写姓名、单位或年级、分工', '待负责人填写联系人、电话、邮箱']
    for row, value in zip(document.tables[0].rows, values):
        row.cells[1].text = value
    for table, text in zip(document.tables[1:6], content['information']):
        table.cell(0, 0).text = text
    document.tables[6].cell(0, 0).text = content['access']
    # Delete only the explanatory hint paragraphs; retain all field names,
    # table order, declaration and the unfilled signature/date line.
    paragraphs = list(document.paragraphs)
    for index in [2, 3, 4, 7, 9, 11, 13, 15, 17]:
        node = paragraphs[index]._element
        node.getparent().remove(node)
    document.paragraphs[0].style = 'Title'
    for style in [document.styles['Normal'], document.styles['Title']]:
        style.font.name = 'WenQuanYi Micro Hei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if style.name == 'Normal' else 19)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'WenQuanYi Micro Hei')
        ppr = style.element.find(qn('w:pPr'))
        if ppr is not None:
            border = ppr.find(qn('w:pBdr'))
            if border is not None: ppr.remove(border)
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'WenQuanYi Micro Hei'
            run.font.color.rgb = RGBColor(0, 0, 0)
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'WenQuanYi Micro Hei')
    for paragraph in document.paragraphs[:2]:
        properties = paragraph._element.find(qn('w:pPr'))
        if properties is not None:
            border = properties.find(qn('w:pBdr'))
            if border is not None: properties.remove(border)
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.18
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs:
                        run.font.name = 'WenQuanYi Micro Hei'
                        run.font.size = Pt(11)
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'WenQuanYi Micro Hei')
                properties = cell._tc.get_or_add_tcPr()
                borders = properties.find(qn('w:tcBorders'))
                if borders is None:
                    borders = OxmlElement('w:tcBorders');properties.append(borders)
                for edge in ['top', 'left', 'bottom', 'right']:
                    e = OxmlElement('w:' + edge)
                    for k, v in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items():e.set(qn('w:' + k), v)
                    borders.append(e)
    document.save(output)
    return {'character_count_upper_bound': counts, 'limits': limits, 'total': sum(counts),
            'needs_human_completion': ['team', 'members', 'contact', 'signature', 'date', 'Word word-count review']}


def make_pdf(output, content, font_path):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.graphics.shapes import Drawing, Rect, Line, String, Polygon
    pdfmetrics.registerFont(TTFont('CJK', str(font_path), subfontIndex=0))
    normal = ParagraphStyle('body', fontName='CJK', fontSize=10.7, leading=16.4, spaceAfter=8, wordWrap='CJK', textColor=colors.black)
    h1 = ParagraphStyle('h1', parent=normal, fontSize=17, leading=23, spaceBefore=0, spaceAfter=14, keepWithNext=True)
    h2 = ParagraphStyle('h2', parent=normal, fontSize=12, leading=18, spaceBefore=8, spaceAfter=6, keepWithNext=True)
    small = ParagraphStyle('small', parent=normal, fontSize=9, leading=13, spaceAfter=4)
    cell_style = ParagraphStyle('cell', parent=normal, fontSize=9.6, leading=14.6, spaceAfter=0)
    summary = json.loads((ROOT/content['evidence_paths']['evaluation']).read_text())['summary']
    class BookmarkedDoc(SimpleDocTemplate):
        def afterFlowable(self, flowable):
            if hasattr(flowable, '_bookmark'):
                key, level = flowable._bookmark
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(flowable.getPlainText(), key, level=level, closed=False)
    doc = BookmarkedDoc(str(output), pagesize=A4, leftMargin=44, rightMargin=44, topMargin=42, bottomMargin=42,
                        title='Claim2Value 复赛项目说明', author='Claim2Value')
    width = A4[0] - 88
    def table(rows, widths):
        cells = [[Paragraph(escape(str(x)), cell_style) for x in row] for row in rows]
        t = Table(cells, colWidths=[width*w for w in widths], repeatRows=1, hAlign='LEFT')
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8EEF2')),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F6F7F8')]),
            ('GRID',(0,0),(-1,-1),0.5,colors.HexColor('#D9D9D9')),
            ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),7),
            ('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        return t
    def diagram():
        d = Drawing(width, 147)
        boxes = [(175,115,155,28,'声明与证据原文'),(155,72,195,28,'来源记录与口径规则'),
                 (14,25,205,28,'自定义核查报告与拒答'),(270,25,223,28,'预置工程经济批判财务链')]
        for x,y,w,h,label in boxes:
            d.add(Rect(x,y,w,h,fillColor=colors.HexColor('#F3F5F6'),strokeColor=colors.HexColor('#A5AEB5'),strokeWidth=0.7))
            d.add(String(x+w/2,y+10,label,fontName='CJK',fontSize=9.5,textAnchor='middle',fillColor=colors.black))
        for x1,y1,x2,y2 in [(252,115,252,100),(225,72,117,53),(282,72,381,53)]:
            d.add(Line(x1,y1,x2,y2,strokeColor=colors.HexColor('#52616B'),strokeWidth=1))
            d.add(Polygon([x2,y2,x2-3,y2+6,x2+3,y2+6],fillColor=colors.HexColor('#52616B'),strokeColor=None))
        d.add(String(117,9,'不绑定估值参数',fontName='CJK',fontSize=9,textAnchor='middle'))
        d.add(String(381,9,'仅限固定公司示例',fontName='CJK',fontSize=9,textAnchor='middle'))
        return d
    story = []
    labels = {'evidence_absence':'证据缺失','qualifier_removal':'限定词删除','source_downgrade':'来源降级','temporal_shift':'时间错位','unit_swap':'口径偷换','value_tampering':'数值篡改'}
    for i, page in enumerate(content['pages']):
        if i:story.append(PageBreak())
        if i == 0:
            story.append(Paragraph('Claim2Value 复赛项目说明',h1))
            story.append(Paragraph('北京大学金融AI智能体创新大赛　版本日期 2026年10月9日',small))
            story.append(Spacer(1,10))
        p = Paragraph(escape(page['heading']),h1)
        if i in (1,3):p._bookmark=('detail'+str(i),1)
        else:p._bookmark=({0:'section1',2:'section2',4:'section3',5:'section5'}[i],0)
        story.append(p)
        for item in page['items']:
            kind = item['kind']
            if kind in ['h1','h2','p']:
                p = Paragraph(escape(item['text']),{'h1':h1,'h2':h2,'p':normal}[kind])
                if 'bookmark' in item:p._bookmark=(item['bookmark'],0)
                story.append(p)
            elif kind=='table':story.extend([table(item['rows'],item['widths']),Spacer(1,9)])
            elif kind=='diagram':story.extend([diagram(),Spacer(1,9)])
            elif kind=='metrics_table':
                rows=[['开发扰动类别','严格标签命中','比例']]
                for k,v in summary['mutations']['by_type'].items():rows.append([labels[k],f"{v['hits']}/{v['n']}",f"{v['rate']:.1%}"])
                rows.append(['合计 19个原家族','87/98','88.8%'])
                story.extend([table(rows,[0.5,0.25,0.25]),Spacer(1,9)])
            elif kind=='scenario_table':
                rows=[['固定示例 EV 亿元','base','upside','downside']]
                for stem,label in [('green_demo','绿的谐波'),('shuanghuan_demo','双环传动')]:
                    f=json.loads((ROOT/content['evidence_paths']['cases']/f'{stem}.json').read_text())['financial']['scenarios']
                    rows.append([label]+[f"{f[k]['enterprise_value_bn']*10:.2f}" for k in ['base','upside','downside']])
                story.extend([table(rows,[0.43,0.19,0.19,0.19]),Spacer(1,9)])
            elif kind=='runtime':
                r=summary['runtime'];env=summary['environment']
                text=f"本轮环境为Linux x86_64、Python {env['python']}，运行环境报告{env['cpu_logical_count']}个逻辑CPU。每条扰动重复{r['repeats_per_mutation']}次，先取各题中位耗时，再对98个中位数汇总：p50为{r['p50_case_median_ms']:.3f}ms，p95为{r['p95_case_median_ms']:.3f}ms。所有重复的判断标签一致。"
                story.append(Paragraph(escape(text),normal))
    def footer(canvas, doc):
        canvas.saveState();canvas.setFont('CJK',8.5);canvas.setFillColor(colors.HexColor('#58626A'))
        canvas.drawString(44,24,'Claim2Value　2026年10月9日');canvas.drawRightString(A4[0]-44,24,str(doc.page));canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--information-template',type=Path,required=True)
    parser.add_argument('--font',type=Path,required=True)
    parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args();args.out.mkdir(parents=True,exist_ok=False)
    content=json.loads((ROOT/'docs/semifinal/materials_content.json').read_text())
    counts=make_information(args.information_template,args.out/'01_统一信息表_草稿.docx',content)
    make_pdf(args.out/'02_项目说明_复赛草稿.pdf',content,args.font)
    (args.out/'草稿说明.json').write_text(json.dumps({'status':'DRAFT_NOT_FINAL_SUBMISSION','information':counts,'next':['new-source verification','independent validation','team fields/signature','trial feedback','presentation/video','actual technical results']},ensure_ascii=False,indent=2))
    print(json.dumps({'output':str(args.out),'information':counts},ensure_ascii=False,indent=2))


if __name__=='__main__':main()

