"""Build a reviewable package from current receipts; never sign or submit for a team."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import shutil
import sys
import zipfile
from pathlib import Path
from decimal import Decimal

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.build_semifinal_materials import make_information, make_pdf
from scripts.build_local_deployment import build as deployment_zip


def read(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def material_content(work, verification):
    summary = read(work / "development_evaluation_utf8/results.json")["summary"]
    runtime = summary["runtime"]
    count = verification["software_tests_passed"]
    information = [
        "Claim2Value面向机器人产业链的产业与投资研究，将企业技术声明、证据原文、口径提示和工程到财务的假设串联起来。可操作作品提供两家公司固定情景、自定义规则核查、历史财务核验和报告导出。缺证或规则不足时拒答；自定义输入不自动估值。历史财务展示34条已核对公开年报原页的合并年度数据，来源与披露日期可追溯；待核CSMAR数据与已核事实分开。实时检索、PDF自动认证和在线模型尚未接入。",
        "本轮新增历史财务核验及导出，接入34条年报原页样本，修正历史输入单位；查收CSMAR核心三表，270项原值分列待核，并建立公告续取与源文件保护。完成Windows页面、部署与技术题独立复现，统一更新材料和实际运行记录。固定估值情景和技术题名单保持冻结。",
        "目标是帮助研究者核对技术叙述是否遗漏测试条件、混用指标或时间口径，避免把技术优势和客户覆盖直接当作收入与估值依据。系统记录输入、原文、出处、指纹和中间状态，通过显式规则检查定义、数值、条件与来源提示；确定性问题给出提示，覆盖不足时拒答。固定示例提供工程、经济映射、替代解释和敏感性情景，自定义输入仅核查。新增财务接收器按真实披露日、元单位、合并/母公司和年度/中期口径检查，未满足条件的数据保留待核。创新体现在可复查的工作流及事实、假设和未解决问题的明确分离。自研部分为规则、证据记录、映射、批判、财务接收和交互；Streamlit等为外部组件。文件哈希不等于来源认证；未开展独立实验来证明优于商业模型。",
        f"本次Windows复测既有开发库98条扰动，来自19个原家族，标签命中87/98、家族等权89.5%；总是拒答诊断基线19/98。开发库已用于规则开发，不能称样本外业务准确率。11条未命中完整保留，10条人工开发边界题命中10/10。{count}项软件检查通过；通过数不代表业务效果。34条公开财务样本来源年报原页，严格接收检查错误0。CSMAR三表45个目标格齐全，但当前正式定义仍缺，270项原值待核。暖启动本地函数每题重复5次，中位耗时p50约{runtime['p50_case_median_ms']:.3f}ms、p95约{runtime['p95_case_median_ms']:.3f}ms，不含页面、人工阅读、网络和模型。API调用与费用均0，算力货币成本未计量。Windows独立部署实际启动及浏览器操作通过。技术题24,240行冻结名单独立解压复现与官方检查通过；2020已查看，扣费Sharpe−0.091268，测试期收益与排名未知。",
        "启动本地页面后可选择两家公司示例、填写声明与证据，或进入历史财务核验；查看来源、规则、期间口径与边界，下载JSON或Markdown。已核年报样本仅覆盖绿的谐波；步科、双环传动显示缺口，本机CSMAR待核区不进入已核事实。固定情景属于假设演示，其中既有2025情景不作为当前预测。空声明和错误链接提示错误，缺证拒答。Windows部署与实际浏览器自动操作已验证，真人试用尚待负责人完成。首次安装依赖需联网，运行不需账号或模型密钥；评委可使用06本地部署。迁移需要新增定义、独立标签、公司原件与参数审核。PDF设书签，技术证据可定位案例、失败记录、财务和复现。"
    ]
    pages = read(ROOT / "docs/semifinal/materials_content.json")["pages"]
    # Retain the reviewed method exposition while replacing all stage-specific claims.
    pages[1]["items"][-1]["text"] = "历史Claim Bank保留47条已验证和4条待验证的旧标注，不将其当作本轮重新认证的成绩。GH_007、BK_003、SH_002、BK_001继续保留版本、工况或一手依据缺口。新CSMAR原值和标题线索不自动升级技术声明。"
    pages[1]["items"][1]["rows"].insert(6, ["历史财务核验", "已实现", "34条公开年报原页样本；披露日、口径和定位可追溯，待核区隔离"])
    pages[2]["items"][-1]["text"] = f"本次Windows独立部署启动与实际浏览器操作通过，具体用时、环境和文件哈希见05_技术证据/部署验收.json。{count}项软件检查通过，属于工程验证。首次安装依赖需联网，后续依赖齐备时离线运行；没有非开发成员试用结果。"
    pages[3]["items"][2]["text"] += " 历史财务纠错不改变冻结情景。2025期是既有假设情景，不作为当前年度预测。"
    pages[4]["items"][0]["text"] = "提供本地离线部署，默认无需账号与模型密钥。Python推荐3.12；Windows双击start_demo.bat，Linux执行bash start.sh。首次安装依赖需联网。本次Windows实际启动与浏览器操作通过，真人试用仍待完成。"
    pages[4]["items"][1]["text"] += " 新增路径：进入历史财务核验，筛选公司与期间，查看真实披露日和原页定位，再导出历史参考报告。"
    pages[4]["items"][4]["text"] = "技术证据索引把能力对应到真实案例、开发评测及部署记录；PDF按评分项设置书签。开发数据和11条失败见05_技术证据/开发复测，运行案例见05_技术证据/实际案例，页面下载见*_actual_download.json。自研代码位于06_本地部署/src；原始商业数据库导出、赛事原始数据、校园账号与密钥不进入本包。"
    pages[4]["items"][7]["text"] = "10月3日厂商网页核验记录中，SHPR-20E与RV-20E在15 rpm下均为167 Nm、4.7 kg，密度35.53 Nm/kg是推算。未保存原件字节，网页版本日期不明，不能认证旧版参数。LHS现版与旧研报不同，未混版计算。四条缺证状态未升级；本次财务接入与Windows复现另有收据，真人试用仍待完成。"
    pages[5]["items"][2]["text"] = "由Codex继续核对公开原件，四条待证声明需要具体版本、工况、PDF页码或原始统计定义；当前缺口如实保留。CSMAR续取入口默认离线预览，只执行剩余公告ID。额度恢复时间未知，本轮不重试已耗尽的账号。"
    pages[5]["items"][5]["text"] = "技术题使用主办方匿名400家公司及100因子，独立于作品公司分析。2020已查看、扣费Sharpe为−0.091268，不能再称未触碰留出。12个候选未达到替换标准，保留原有效名单。本机恢复官方测试文件后，以锁定依赖解压当前code.zip复现，CSV哈希一致、官方检查有效。测试收益不可见，不报告比赛Sharpe或排名。"
    pages[5]["items"][7]["text"] = "队名与成员姓名、单位按负责人提供的信息填写。尚需负责人补齐成员分工、联系人、签名及日期，完成一次真人试用、最终材料人工审阅及赛事上传回执。本包尚未正式提交。"
    pages[5]["items"][9]["text"] = "赛事结构依据附件一、附件二及复赛通知，原件在赛事材料Drive目录。实现和复跑说明见github.com/aeiou0123/pku-financial-ai-agent。本次结果、失败和环境见05_技术证据/开发复测/results.json；案例与固定情景数字见实际案例/。技术题Windows复现见技术题本机复现.json，视频格式和完整解码见视频验收.json。旧云端记录与本次运行分别登记；仍未恢复的历史原文保留缺口。"
    pages[5]["items"][10]["text"] = "厂商历史参数来源与栏目定位见05_技术证据/primary_source_review_20261003/review.md及review.json。环动来源finemotion.com.cn/home/Index/product?id=34；纳博特斯克来源precision.nabtesco.com/tw/products/detail/RV-E；绿的来源leaderdrive.com/product/4.html。保留获取日期与历史版本限制，记录文件指纹不作为原件指纹；本次未将四条待核技术声明升级。"
    financial = {"heading": "2 历史财务证据与待核边界", "items": [
        {"kind": "p", "text": "绿的谐波2025年报中的2024、2025合并报表，17项指标各两年，共34条。接收器保持16列接口，按Decimal换算金额、核对披露日、原文件哈希及报表和期间口径。两年资产－负债－总权益差额均0.00元。工具的checked状态仅表示已执行数值检查通过，原页核验另有记录。"},
        {"kind": "table", "widths": [0.39, 0.36, 0.25], "rows": [["2024公开锚点", "原值 元", "来源"],
            ["营业收入", "387,411,303.84", "2024年报PDF第7页"], ["归母净利润", "56,168,149.88", "同页"],
            ["经营净现金流", "27,981,461.73", "同页"], ["总资产", "3,755,317,295.17", "同页"],
            ["归母权益", "3,425,332,272.82", "同页"]]},
        {"kind": "p", "text": "2024年报披露于2025-04-30；34条比较样本实际来自2025年报，披露于2026-04-23。其中2024比较值不能冒充2024实时已知数据。收入和归母净利润历史输入由误写的3.87、0.561681十亿元修正为0.38741130384、0.05616814988，其他假设与情景结果保持不变。"},
        {"kind": "p", "text": "CSMAR核心三表45/45目标格齐全，270项待核原值逐值保留字段代码、客户端包装JSON定位和哈希。15个合并观察的45项勾稽中42项精确为0、1项0.01元尾差、2项因少数权益缺失未检查。当前正式定义、披露日期及中期累计口径不完整，不能据此生成经过认证的财务预测；中期不差分、不年化。"},
        {"kind": "p", "text": "目标公告还缺1,333个ID，整理为9个最多200个ID的续取任务。基金记录隔离，旧原件逐文件保护，新增文件单独登记。当前实际额度已耗尽，恢复时刻未知；本轮只完成离线预览。部署包不含商业原始导出、赛事原始数据或账号凭据。"}
    ]}
    pages.insert(4, financial)
    return {"date": "2026-10-10", "information": information,
            "access": "06_本地部署：Python3.12推荐，Windows双击start_demo.bat；Linux执行bash start.sh。首次安装依赖需联网，后续依赖齐备时离线运行。浏览器http://127.0.0.1:8501，默认无需账号、无需模型密钥。Windows已实际验收，真人试用仍待完成。",
            "pages": pages, "evidence_paths": {"evaluation": str((work / "development_evaluation_utf8/results.json").relative_to(ROOT)),
                         "cases": str((work / "actual_cases").relative_to(ROOT))}}


def slides(path, panels):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5)
    def text(s, value, x, y, w, h, size=22, bold=False, color="18232D"):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame; tf.word_wrap = True
        for i, line in enumerate(value.split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.text = line
            para.space_after = Pt(12)
            for run in para.runs: run.font.name = "Microsoft YaHei"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(color)
        return box
    for index, panel in enumerate(panels, 1):
        s = p.slides.add_slide(p.slide_layouts[6])
        text(s, panel["title"], .65, .38, 12, 1, 31, True, "102E42")
        if panel.get("table"):
            rows = panel["table"]; shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(.65), Inches(1.7), Inches(12), Inches(3.6))
            t = shape.table
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    c = t.cell(ri, ci); c.text = str(value); c.fill.solid(); c.fill.fore_color.rgb = RGBColor.from_string("E8EEF2" if ri == 0 else "F6F7F8")
                    for para in c.text_frame.paragraphs:
                        for run in para.runs:run.font.name = "Microsoft YaHei"; run.font.size = Pt(20); run.font.color.rgb = RGBColor.from_string("18232D")
            if panel.get("body"):text(s, panel["body"], .65, 5.6, 12, 1, 21, color="58626A")
        else:text(s, panel["body"], .65, 1.65, 12, 4.75, 27 if len(panel["body"]) < 180 else 24)
        text(s, f"Claim2Value · 2026-10-10                        {index}", .65, 6.97, 12, .3, 12, color="58626A")
        s.notes_slide.notes_text_frame.text = panel.get("notes", panel.get("body", ""))
    p.save(path)


def main_panels(work, verification):
    summary = read(work / "development_evaluation_utf8/results.json")["summary"]
    scenarios = []
    for stem, company in [("green_demo", "绿的谐波"), ("shuanghuan_demo", "双环传动")]:
        f = read(work / f"actual_cases/{stem}.json")["financial"]["scenarios"]
        scenarios.append([company] + [f"{f[k]['enterprise_value_bn']*10:.2f}" for k in ("base", "upside", "downside")])
    return [
        {"title": "Claim2Value", "body": "机器人产业链声明核查与历史财务证据\n\n北京大学金融AI智能体创新大赛复赛\n本地规则模式 · 事实与假设分开呈现"},
        {"title": "研究者需要可复查的口径和依据", "body": "减重30%需要保留同等出力与比较对象。\n额定与峰值、归母与扣非、年度与中期分别保留口径。\n客户覆盖不自动证明订单、收入或利润兑现。\n来源文件与摘录能被复查，缺口能被看见。"},
        {"title": "工作流与自研边界", "table": [["环节", "实现与限定"], ["输入与证据", "保留声明、原文、出处与指纹"], ["自定义核查", "规则提示与拒答，不自动估值"], ["固定示例", "两家公司工程、经济映射、批判与情景"], ["历史财务", "34条公开年报样本，待核数据隔离"], ["外部组件", "Streamlit、openpyxl、NumPy/pandas/PyArrow"]], "body": "来源等级与规则分数是提示，不认证原件或校准概率。"},
        {"title": "历史财务原页核验", "table": [["2024指标", "原值 元"], ["营业收入", "387,411,303.84"], ["归母净利润", "56,168,149.88"], ["经营净现金流", "27,981,461.73"], ["总资产", "3,755,317,295.17"], ["归母权益", "3,425,332,272.82"]], "body": "2024年报PDF第7页。34条比较样本来自2025年报，披露日2026-04-23。"},
        {"title": "已核事实与CSMAR待核区", "body": "核心三表45/45目标格已收到，270项原值带字段、定位和哈希。\n当前正式定义、披露日和中期口径未补齐，保持待核。\n45项勾稽：42项精确为0，1项0.01元尾差，2项因缺值未检查。\n公告尚缺1,333个目标A股ID，9批续取；额度恢复时刻未知。\n商业原始导出及凭据不进入部署包。"},
        {"title": "开发复测与失败记录", "table": [["验证", "结果", "解释"], ["既有开发扰动", "87/98", "19个原家族，非样本外"], ["家族等权", "89.5%", "规则曾使用该库开发"], ["拒答诊断基线", "19/98", "不能证明优于商业模型"], ["人工开发边界题", "10/10", "不是独立留出"], ["软件检查", str(verification['software_tests_passed']) + "通过", "工程检查，非业务效果"]], "body": "11条未命中完整保留；自然来源与业务准确率仍需独立验证。"},
        {"title": "保留固定假设情景", "table": [["EV 亿元", "base", "upside", "downside"], *scenarios], "body": "不是股价、当前预测或因果效应。2025期为既有假设，历史纠错不自动改情景。"},
        {"title": "本次Windows真实运行", "body": f"历史财务页面、公司筛选、缺证提示与报告下载可实际操作。\n独立部署通过；浏览器自动操作与真人试用分别记录。\n暖启动函数每题重复5次：p50 {summary['runtime']['p50_case_median_ms']:.3f}ms，p95 {summary['runtime']['p95_case_median_ms']:.3f}ms。\n上述耗时不含页面、阅读、网络和模型。API调用0，算力费用未计量。"},
        {"title": "技术题保持冻结方案", "body": "400家匿名公司、100因子，与作品财务分析独立。\n1,212个测试日，每日20家公司，共24,240行。\n官方测试文件已恢复并核对哈希。\n独立解压复现CSV一致，官方检查通过。\n2020已查看，扣费Sharpe -0.091268；测试收益及排名未知。"},
        {"title": "评委使用路径", "body": "06_本地部署：Windows双击start_demo.bat。\n默认无需账号或密钥，首次安装依赖需联网。\n选择固定示例、自定义核查或历史财务核验，再下载报告。\n视频保存连续原始录屏与实际下载文件，配四部分章节。"},
        {"title": "边界与负责人收尾", "body": "没有实时检索、PDF自动认证或当前在线大模型。\n四条技术声明及CSMAR完整定义仍有缺口，不自动升级。\n没有独立业务效果或真人试用结论。\n负责人补成员分工、联系人、签名、日期，完成试用与人工审阅。\n官方截止：2026-10-13 00:00 北京时间；上传后保存回执。"}
    ]


def build(args):
    work = args.work.resolve(); verification = read(args.verification)
    if verification["technical_reproduction"] != "PASS" or verification["windows_deployment"] != "PASS":
        raise ValueError("Current reproduction and Windows deployment must pass before material generation")
    out = args.out.resolve(); out.mkdir(parents=True, exist_ok=False)
    evidence = out / "05_技术证据"; evidence.mkdir()
    content = material_content(work, verification)
    if args.team_information:
        content['team_information'] = read(args.team_information)
    (work / "materials_content_current.json").write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8")
    counts = make_information(args.template, out / "01_统一信息表.docx", content)
    # Use a font actually installed on the reviewer platform.
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(out / "01_统一信息表.docx")
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = "Microsoft YaHei"; style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    for container in [*doc.paragraphs, *[p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]]:
        for run in container.runs:run.font.name = "Microsoft YaHei"; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.save(out / "01_统一信息表.docx")
    make_pdf(out / "02_项目说明.pdf", content, args.font)
    slides(out / "03_展示稿.pptx", main_panels(work, verification))
    shutil.copy2(args.video, out / "04_运行视频.mp4")
    technical = out / "07_技术题"; shutil.copytree(work / "frozen_technical", technical)
    # The strategy and original report are frozen; append a dated reproducibility record only.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    report = Presentation(technical / "技术题报告.pptx")
    for original_slide in report.slides:
        for shape in original_slide.shapes:
            if not shape.has_text_frame:continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace('Windows与跨系统一致性未实测。', '10月3日仅同环境复现；Windows补记见第14页。')
                    if run.text == '实测运行时间与环境':run.text = '10月3日云端运行时间与环境'
    layout = min(report.slide_layouts, key=lambda l: len(l.placeholders))
    s = report.slides.add_slide(layout)
    for placeholder in list(s.placeholders):
        placeholder._element.getparent().remove(placeholder._element)
    box = s.shapes.add_textbox(Inches(.6), Inches(.5), report.slide_width - Inches(1.2), report.slide_height - Inches(1))
    box.text_frame.word_wrap = True
    box.text_frame.text = "2026-10-10 Windows独立复现补记\n\n已恢复官方test.parquet，144,652,201字节，SHA-256与官方一致。\n锁定依赖解压当前code.zip运行reproduce.py，官方check有效。\n1,212日、每天20股、24,240行，CSV SHA-256与原冻结名单一致。\n保留原模型与参数；2020已经查看，测试收益与排名未知。\n旧运行收据与当前文档更新版code.zip分别登记，ZIP哈希不混同。"
    for para in box.text_frame.paragraphs:
        para.space_after = Pt(18)
        for run in para.runs:run.font.name = "Microsoft YaHei"; run.font.size = Pt(23)
    if len(report.slides) > 20:raise ValueError("Technical report exceeds 20 slides")
    report.save(technical / "技术题报告.pptx")
    dep = work / "final_deployment.zip"; deployment_zip(dep, work / "deployment_receipt.json")
    with zipfile.ZipFile(dep) as z:z.extractall(out)
    copies = [(work / "verification.json", "本次验收.json"), (work / "deployment_receipt.json", "部署验收.json"),
              (work / "technical_reproduction_receipt.json", "技术题本机复现.json"),
              (work / "financial_integration/historical_public_financial.json", "历史财务公开样本.json"),
              (work / "video/recording_receipt.json", "录屏记录.json")]
    copies += [(work / "technical_historical_binding.json", "技术题历史收据绑定.json"),
               (work / "source_integrity_after.json", "旧原件保护.json"),
               (work / "intake_separation_check.json", "财务待核隔离检查.json"),
               (work / "resume_preview_public_summary.json", "公告续取离线预览.json"),
               (work / "software_tests_final.xml", "软件检查.xml"),
               (work / "software_tests_final.log", "软件检查.log"),
               (work / "video_validation.json", "视频验收.json")]
    for source, name in copies:shutil.copy2(source, evidence / name)
    shutil.copytree(work / "development_evaluation_utf8", evidence / "开发复测")
    shutil.copytree(work / "actual_cases", evidence / "实际案例")
    shutil.copytree(ROOT / "docs/semifinal/evidence/primary_source_review_20261003", evidence / "primary_source_review_20261003")
    for path in (work / "video").glob("*_actual_download.json"):shutil.copy2(path, evidence / path.name)
    (evidence / "信息表限字预检.json").write_text(json.dumps(counts, ensure_ascii=False, indent=2), encoding="utf-8")
    (evidence / "证据索引.md").write_text("# 本次提交证据索引\n\n"
        "- 本次验收.json：本次软件检查、Windows部署和技术题复现。\n"
        "- 历史财务公开样本.json：34条年报原页核对值、真实披露日、原件SHA和定位；原始PDF不再分发，URL在来源记录中。\n"
        "- 开发复测/：98条既有开发题、19个家族、11条失败；不是独立业务准确率。\n"
        "- 实际案例/与*_actual_download.json：本次运行和视频实际页面下载，含失败及拒答案例。\n"
        "- 录屏记录.json：连续真实页面录制，浏览器自动操作，非真人试用；原始WebM留在工作区。\n"
        "- 部署验收.json：从独立目录启动Windows一键入口并操作页面。代码见../06_本地部署/。\n"
        "- 技术题本机复现.json：当前code.zip与submission.csv哈希；旧云端收据的ZIP哈希另列。\n"
        "- CSMAR商业原始导出留在授权私有工作区，270项原值仍待正式定义与披露日；不作为已核事实。\n"
        "- 对话/云端Work覆盖清单保留原文、重建与未获得状态；没有宣称全部历史原文已恢复。\n\n"
        "软件通过与来源文件哈希不代表声明、投资效果或因果效应已得到验证。\n", encoding="utf-8")
    (evidence / "负责人收尾.md").write_text("# 正式上传前必须完成\n\n"
        "1. 已按负责人信息填写队名和三位成员姓名、单位；补齐成员分工、联系人电话邮箱、签名与日期。\n"
        "2. 让一位非开发成员实际启动、操作、导出报告，记录真实反馈；不能把自动化录屏当真人试用。\n"
        "3. 人工查看信息表、PDF、两份PPT和完整视频，核对同一案例与能力边界。\n"
        "4. 总包正式文件名为饮水思源_Claim2Value.zip，按赛事指定入口上传并保存回执。\n"
        "5. 官方截止2026-10-13 00:00北京时间；以赛事后续公告为准。本包尚未正式上传。\n", encoding="utf-8")
    return {"output": str(out), "information": counts, "main_slides": 11, "technical_slides": len(report.slides),
            "status": "REVIEW_PACKAGE_PENDING_TEAM_SIGNATURE_HUMAN_TRIAL_UPLOAD"}


if __name__ == "__main__":
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--work", type=Path, required=True); p.add_argument("--out", type=Path, required=True)
    p.add_argument("--verification", type=Path, required=True); p.add_argument("--template", type=Path, required=True)
    p.add_argument("--font", type=Path, required=True); p.add_argument("--video", type=Path, required=True)
    p.add_argument("--team-information", type=Path)
    a=p.parse_args(); print(json.dumps(build(a), ensure_ascii=False, indent=2))
