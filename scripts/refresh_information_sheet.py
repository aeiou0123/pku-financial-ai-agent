# -*- coding: utf-8 -*-
"""Refresh the semifinal information sheet from docs/semifinal/materials_content.json.

Keeps the official template layout (7 tables, field names, signature line) and
replaces only the content cells. Never invents team fields or signatures.
"""
from __future__ import annotations
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "deliverables/semifinal_20261012/饮水思源_Claim2Value/01_统一信息表.docx"
LIMITS = [300, 200, 800, 700, 500]


def set_cell(cell, text):
    """Replace a cell's text, keeping the first run's formatting."""
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def main():
    from docx import Document
    content = json.loads((ROOT / "docs/semifinal/materials_content.json").read_text(encoding="utf-8"))
    counts = [len("".join(t.split())) for t in content["information"]]
    if any(n > lim for n, lim in zip(counts, LIMITS)) or sum(counts) > 2500:
        raise SystemExit(f"Information fields exceed limits: {counts}")

    document = Document(str(TARGET))
    if len(document.tables) != 7:
        raise SystemExit(f"Unexpected table count: {len(document.tables)}")

    identity = content.get("team_information", {})
    values = [identity.get("team_name", "饮水思源"), "Claim2Value",
              identity.get("members_text", "魏孟祺 — 西安交通大学；陈洛笛 — 西安交通大学；孙圣尧 — 上海交通大学\n（各人分工待负责人填写）"),
              identity.get("contact_text", "联系人姓名、电话、邮箱：待负责人填写")]
    for row, value in zip(document.tables[0].rows, values):
        set_cell(row.cells[1], value)
    for table, text in zip(document.tables[1:6], content["information"]):
        set_cell(table.cell(0, 0), text)
    set_cell(document.tables[6].cell(0, 0), content["access"])
    document.save(str(TARGET))
    print(json.dumps({"counts": counts, "limits": LIMITS, "total": sum(counts)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
